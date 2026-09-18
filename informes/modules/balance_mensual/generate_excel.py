"""
Arma el Excel de Balance Mensual (1/7/25-31/8/26) que pidio Javier, con
detalle de cliente/proveedor en una segunda hoja.

Fuentes (ya exportadas con export.py de este modulo y modules/iibb/export.py):
- exploracion/balance_mensual_export.csv: "Reporte Balance Mensual"
  (Consultas > Contabilidad > Balance Mensual), agregado por cuenta x mes.
  NO tiene columna de cliente/proveedor -- es estructuralmente imposible
  sacarle ese detalle (confirmado en vivo, 2026-09-16).
- exploracion/iibb_export.xlsx: export crudo de "Imputaciones" (Consultas >
  Contabilidad > Imputaciones, TODAS las cuentas sin filtrar -- el filtro a
  cuentas de IIBB pasa recien en modules/iibb/ingest.py, este archivo es el
  export plano). La columna "Leyenda" trae el nombre real de cliente o
  proveedor por linea de asiento (confirmado en vivo: p.ej. cuenta 112110
  DEUDORES EN CTAS. CTES. con Leyenda "ALUAR ALUMINIO ARGENTINO...", o
  cuenta 211020 PROVEEDORES DE PRODUCCION con Leyenda "LOPETEGUI JUAN
  MANUEL"). Es texto libre, no una columna de entidad limpia y garantizada
  en el 100% de las lineas -- Javier lo pidio asi ("La base de
  imputaciones?", confirmado 2026-09-16) sabiendo esa limitacion.

Uso:
    python -m modules.balance_mensual.export
    python -m modules.iibb.export
    python -m modules.balance_mensual.generate_excel
"""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows

from common import normalizar_numero

PERIODO_DESDE = (2025, 7)   # (año, mes) inclusive
PERIODO_HASTA = (2026, 8)   # (año, mes) inclusive

BALANCE_CSV = Path("exploracion/balance_mensual_export.csv")
IMPUTACIONES_XLSX = Path("exploracion/iibb_export.xlsx")
SALIDA = Path("salida/balance_mensual_2025-07_2026-08.xlsx")

COLUMNAS_MONTO_BALANCE = ["Saldo Anterior", "Debe", "Haber", "Neto Mes", "Saldo Actual"]
COLUMNAS_DETALLE_IMPUTACIONES = [
    "Cuenta", "Nombre Cuenta", "Sub.Cta.", "Fecha", "Asiento", "TA", "TR",
    "N° Referencia", "Moneda", "Centro Costo", "Importe", "Leyenda",
]


def cargar_balance_mensual() -> pd.DataFrame:
    df = pd.read_csv(BALANCE_CSV, encoding="utf-8", sep=";", engine="python")
    for col in COLUMNAS_MONTO_BALANCE:
        df[col] = df[col].apply(normalizar_numero)
    periodo = df["Año"] * 100 + df["Mes"]
    desde = PERIODO_DESDE[0] * 100 + PERIODO_DESDE[1]
    hasta = PERIODO_HASTA[0] * 100 + PERIODO_HASTA[1]
    df = df[(periodo >= desde) & (periodo <= hasta)].copy()
    return df.sort_values(["Año", "Mes", "Nombre"]).reset_index(drop=True)


def cargar_detalle_imputaciones() -> pd.DataFrame:
    df = pd.read_excel(IMPUTACIONES_XLSX)
    df.columns = [c.replace("N� Referencia", "N° Referencia") if "Referencia" in c else c for c in df.columns]
    df["Fecha"] = pd.to_datetime(df["Fecha"], errors="coerce")
    desde = pd.Timestamp(PERIODO_DESDE[0], PERIODO_DESDE[1], 1)
    hasta_mes = pd.Timestamp(PERIODO_HASTA[0], PERIODO_HASTA[1], 1) + pd.offsets.MonthEnd(0)
    df = df[(df["Fecha"] >= desde) & (df["Fecha"] <= hasta_mes)].copy()
    columnas = [c for c in COLUMNAS_DETALLE_IMPUTACIONES if c in df.columns]
    df = df[columnas]
    return df.sort_values(["Cuenta", "Fecha"]).reset_index(drop=True)


def escribir_hoja(ws, df: pd.DataFrame, titulo: str):
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for col_cells in ws.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(largo + 2, 40)


def main():
    if not BALANCE_CSV.exists():
        print(f"Falta {BALANCE_CSV}. Corre 'python -m modules.balance_mensual.export' primero.")
        return
    if not IMPUTACIONES_XLSX.exists():
        print(f"Falta {IMPUTACIONES_XLSX}. Corre 'python -m modules.iibb.export' primero.")
        return

    balance = cargar_balance_mensual()
    detalle = cargar_detalle_imputaciones()

    wb = Workbook()

    ws0 = wb.active
    ws0.title = "Resumen"
    ws0.append(["Balance Mensual 1/7/2025 - 31/8/2026 - ALESTE ADS S.A."])
    ws0["A1"].font = Font(bold=True, size=14)
    ws0.append([f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    ws0.append([])
    ws0.append(["Hoja 'Balance Mensual': saldo por cuenta contable y mes (Consultas > Contabilidad > Balance Mensual)."])
    ws0.append([f"  Cuentas x mes en el rango: {len(balance)}"])
    ws0.append([])
    ws0.append(["Hoja 'Detalle Cliente-Proveedor': movimientos linea por linea (Consultas > Contabilidad > Imputaciones),"])
    ws0.append(["  con la columna 'Leyenda' como cliente/proveedor de cada linea. Es texto libre de Advertys,"])
    ws0.append(["  no una columna de entidad garantizada en el 100% de las filas -- revisar caso a caso."])
    ws0.append([f"  Lineas en el rango: {len(detalle)}"])

    ws1 = wb.create_sheet("Balance Mensual")
    escribir_hoja(ws1, balance, "Balance Mensual por cuenta (1/7/25-31/8/26)")

    ws2 = wb.create_sheet("Detalle Cliente-Proveedor")
    escribir_hoja(ws2, detalle, "Detalle de Imputaciones con Cliente/Proveedor (Leyenda) - 1/7/25-31/8/26")

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"OK: excel generado en {SALIDA}")


if __name__ == "__main__":
    main()
